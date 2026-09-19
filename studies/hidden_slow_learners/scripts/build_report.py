#!/usr/bin/env python3
"""Generate the separate editable findings report from archived study outputs."""
from pathlib import Path
import json
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def cell_shade(cell,fill):
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(shd)

def keep_row(row):
    element=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(element)

def main():
    s=json.loads((ROOT/'results/summary.json').read_text())
    w=json.loads((ROOT/'results/readable_witness.json').read_text())
    metrics=list(csv.DictReader((ROOT/'results/comparison_summary.csv').open()))
    def metric(h,m,key,n=0):
        return float(next(r[key] for r in metrics if int(r['H'])==h and r['method']==m and int(r['noise_id'])==n))
    doc=Document();sec=doc.sections[0]
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
    for name in ['Normal','Body Text']:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(10.5)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.09
    for name,size in [('Title',28),('Heading 1',20),('Heading 2',13),('Heading 3',11)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('173D55')
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
    hp=sec.header.paragraphs[0];hp.text='MACHINE MARKETS  /  RESEARCH COMPANION 01'
    hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string('526777')
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=foot.add_run('18 September 2026  |  ');run.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    doc.core_properties.title='Hidden Slow Learners: Research Findings and Next Steps'
    doc.core_properties.subject='Executed synthetic study, evidence boundaries, and reproduction roadmap'
    doc.core_properties.author=''
    doc.core_properties.keywords='adaptive markets, stability, reproducibility, hidden learning rates'
    def para(t,bold=False):
        p=doc.add_paragraph();r=p.add_run(t);r.bold=bold;return p
    def heading(t):doc.add_heading(t,2)
    def table(headers,rows,widths=None):
        tb=doc.add_table(rows=1, cols=len(headers));tb.style='Table Grid'
        tb.autofit=False
        for i,h in enumerate(headers):
            tb.rows[0].cells[i].text=h;cell_shade(tb.rows[0].cells[i],'173D55')
            for r in tb.rows[0].cells[i].paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        rep=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(rep)
        for j,row in enumerate(rows):
            cells=tb.add_row().cells
            for i,value in enumerate(row):
                cells[i].text=str(value)
                if j%2==0:cell_shade(cells[i],'F0F4F6')
                for p in cells[i].paragraphs:
                    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(4)
                    for r in p.runs:r.font.size=Pt(9.3)
            keep_row(tb.rows[-1])
        if widths:
            for row in tb.rows:
                for i,wi in enumerate(widths):row.cells[i].width=Inches(wi)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return tb
    def page(title):doc.add_page_break();doc.add_heading(title,1)
    doc.add_heading('Hidden Slow Learners',0)
    para('Research findings and next steps',True)
    para('A separate, executed study supporting the broader “From Human Markets to Machine Markets” research program.')
    heading('The result worth building on')
    para('In the specified mathematical market, a stable system and an unstable system can produce exactly the same early price responses. The hidden difference is how financial decision states update across different learning timescales. A short observation window can therefore support false reassurance even when its measurements are exact.',True)
    para('We constructed these pairs, derived a rate-aware diagnostic, ran the comparisons, and packaged the code and measurements. This is new work within this project. Global novelty and independent external replication have not been established.')
    table(['Completed in this environment','Evidence'],[
        ('Exact paired construction','112 pairs; 224 model instances; opposite stability in every pair'),
        ('Common-shock check','336 paired finite input-response comparisons'),
        ('Measurement experiment','108 configurations; 21,600 independent noisy measurement datasets'),
        ('Diagnostic evaluations','86,400 rows across four declared horizons; six reported methods'),
        ('Verification','29 tests; separate rational and measurement verifiers; clean rerun')],[2.65,4.1])
    heading('The honest status')
    para('This is a complete reproducible research artifact for a restricted positive linear model. It is not yet a real-market discovery, a trained-AI trading study, a journal-ready claim of first discovery, or a demonstrated paradigm shift. The important progress is moving from a broad idea to a precise theorem, an explicit counterexample, executable evidence, and an audit trail.')

    page('1  |  The exact scientific finding')
    heading('The question')
    para('Can a finite stress-test window distinguish stability when the number and learning rates of hidden adaptive states are not known?')
    heading('The model, in ordinary language')
    para('Price moves depend on ordinary price persistence plus the combined pressure of adaptive forecasting states. Each state updates exponentially at its own fixed rate. The model has constant price impact and nonnegative feedback sensitivities. It omits inventories, leverage, spreads, strategic optimization, and changing model parameters.')
    para('Its stability boundary is a + Q = 1, where a is direct price persistence and Q is total cumulative adaptive feedback. This criterion is established positive-system mathematics specialized to the model, not a newly discovered financial law.')
    table(['Quantity','Stable market','Unstable market'],[
        ('Direct persistence a','0.400000','0.400000'),
        ('Active agents after matched splitting','3','3'),
        ('Total hidden feedback Q',f"{w['true_Q_minus']:.6f}",f"{w['true_Q_plus']:.6f}"),
        ('Long-run growth indicator: spectral radius',f"{w['rho_minus']:.6f}",f"{w['rho_plus']:.6f}"),
        ('First four impulse coefficients','1; 0.4; 0.252308; 0.184','Exactly the same'),
        ('Fifth coefficient','0.138428','0.159198')],[2.5,2.1,2.15])
    para('A spectral radius below one means disturbances decay in this fixed linear system; above one means at least one mode amplifies. These numbers do not represent a crisis probability.')
    heading('Why this is stronger than a similar-looking graph')
    para('The construction proves exact equality for every chosen finite horizon, with rational parameters. It also holds for the finite response to any common input sequence using those coefficients. A classifier forced to label one of the two matched possibilities must be wrong on half of a balanced pair. An inconclusive answer avoids that forced mistake.')
    para('The construction uses classical interpolation and truncated-moment identities. It does not refute conventional system identification with a known finite order and sufficient informative data. Some long-horizon constructed pairs lie very close to the stability boundary.')

    page('2  |  What the experiments actually showed')
    heading('Exact checks')
    para('All 112 pairs had opposite stability. The largest floating-point discrepancy between matched price coefficients was 3.33 × 10⁻¹⁶. A separate checker using exact rational arithmetic verified the identities for all pairs and the readable example. Numerical agreement supports the implementation; the proof and rational identities support exact equality.')
    heading('Noisy controlled measurements')
    para('The diagnostic experiment used a stronger information set: controlled measurements of the adaptive feedback kernel, plus a valid externally supplied minimum learning rate. It did not infer the kernel directly from noisy stock returns. There were 100 observation-noise repetitions for each of 108 configurations and two noise levels.')
    rows=[]
    for H in [4,16,64,256]:
        rows.append((H,f"{100*metric(H,'rate_aware','decision_fraction'):.2f}%",'0.00%',f"{100*metric(H,'truncation','wrong_stable_fraction'):.2f}%",f"{100*metric(H,'one_rate','wrong_stable_fraction'):.2f}%"))
    table(['Horizon','Rate-aware: decides','Rate-aware: false stable','Truncation: false stable','One-rate: false stable'],rows,[.65,1.45,1.55,1.55,1.55])
    para('Lower-noise condition. “Decides” is a percentage of all 10,800 datasets. “False stable” is a percentage of the 5,400 truly unstable datasets, not overall error. Different denominators must not be mixed.')
    para('At horizon 256, rate-aware decision coverage was 82.69% under lower noise and 66.20% under higher noise. Across the full experiment, no incorrect confident rate-aware labels occurred. Many cases remained inconclusive; zero errors must always be reported together with this decision coverage.')
    heading('What the conventional comparator does and does not show')
    para('The simple matrix-pencil baseline used positivity checks and the same rate floor, but only up to 32 early taps. At horizon 256 and lower noise it labeled 29.87% of unstable cases stable while deciding on 99.99% of all cases. This is not a fair full-horizon state-of-the-art comparison and does not establish universal superiority. A stronger, equally informed comparator remains a required gate.')
    para('The 128-replicate Gaussian sample means were sampled from their known distribution. The archive does not claim that 2.76 million separate market episodes or trained-agent runs were executed.')

    page('3  |  Assumptions, failures, and decision limits')
    heading('What makes the diagnostic valid')
    table(['Required information','Why it matters'],[
        ('A fixed positive linear model','The total-feedback threshold is not valid for every nonlinear or signed market.'),
        ('Known direct persistence and normalization','An incorrect a or response scale moves the stability boundary.'),
        ('Controlled loop-open kernel observations','A price-only observation is not the same measurement; feedback must be isolated in a trusted replay.'),
        ('A valid minimum active learning rate','An omitted slower state can carry feedback outside the assumed tail bound.'),
        ('Known Gaussian tap-noise standard errors','The finite-family confidence argument relies on this noise model.'),
        ('Four declared observation horizons','The guarantee is not an unlimited, anytime-valid monitoring guarantee.')],[2.45,4.3])
    heading('A failure that must remain visible')
    para('An unstable model with a = 0.4, Q = 0.8, and actual learning rate 0.001 is inconclusive after four exact kernel observations when the correct floor is supplied. Falsely assuming the floor is 0.2 produces a false “stable” label. More precision in those same short measurements does not fix the incorrect architectural assumption.',True)
    para('A separate negative-feedback example also shows that the positive-model threshold cannot be transferred to signed strategies. The code rejects such inputs in its supported model interface and tests the out-of-class example separately.')
    heading('Do not translate these results into trading instructions')
    para('“Abstention” here means declining to classify stability, not ordering agents to stop trading. We have not shown that synchronized trading abstention improves liquidity or reduces financial losses. Likewise, a spectral radius is not a calibrated probability of an actual financial crisis.')
    para('The confidence statement bounds the unconditional chance of any wrong confident label across the four horizons under the assumptions. It does not give a 95% guarantee conditional on the subset of cases that received a decision.')

    page('4  |  Novelty audit and publication positioning')
    heading('What existing work already explains')
    table(['Primary source','Established ingredient / implication'],[
        ('Georges & Pereira (2021)','Adaptive machine-learning agents and market instability are already studied.'),
        ('Perdomo et al. (2020)','Predictions can alter the outcomes they predict; the feedback idea is established.'),
        ('Rantzer (2015)','Positive-system stability and gain analysis supply the mathematical foundation.'),
        ('Oymak & Ozay (2019 preprint)','Finite observations, Markov parameters, and realization already have identification theory.'),
        ('Umenberger & Manchester (2016)','Stable positive-system identification is established; a stable fitted model alone is not evidence of stable truth.'),
        ('Fritzsche, Kirstein & Mädler (2020)','Finite moment sequences can correspond to multiple underlying measures.')],[2.4,4.35])
    para('Full titles, versions, links, and the detailed audit are in literature/LITERATURE_AUDIT.md and the research paper bibliography. The audit checked primary records and selected accessible full texts; it was not a complete systematic priority review.')
    heading('The candidate contribution')
    para('An explicit opposite-stability pair for the adaptive-market model; equal short response and equal active-agent count; a precise distinction between price observations and controlled feedback measurements; a conservative rate-aware diagnostic; and a fully auditable experiment exposing false reassurance and assumption failure.')
    heading('What can be said now')
    para('“We developed a model-specific, reproducible finite-horizon identification result and an assumption-aware diagnostic.”',True)
    heading('What cannot yet be claimed')
    para('“We discovered a new fundamental law of finance,” “the method predicts real crises,” “it is globally first,” “it outperforms the best existing methods,” or “independent researchers reproduced it.” If external review finds the construction already known, the honest options are a replication/translation benchmark or a genuinely new extension—not a new name for the same theorem.')

    page('5  |  Reproduction and what is inside the folder')
    heading('Internal reproducibility completed')
    para('The main computation, an exact-rational checker, and a separately written measurement checker agreed. The measurement checker regenerated every one of the 21,600 archived datasets and recomputed all 86,400 diagnostic decisions without importing the main inference functions. A clean-directory rerun reproduced 17 data files. All checks used the same authoring process and environment; external replication is still open.')
    table(['Folder / file','Purpose'],[
        ('paper/','Complete new LaTeX paper and compiled PDF; proofs and actual results.'),
        ('src/ and tests/','Model, exact construction, bounds, baselines, and 29 tests.'),
        ('configs/main.json','Fixed complete numerical design and seed.'),
        ('scripts/','Experiment runner, two separate verifiers, figures, and report authoring.'),
        ('results/','Rational witnesses, configuration records, measurement inputs, decisions, summaries, and verification logs.'),
        ('docs/ and literature/','Protocol, claim ledger, data dictionary, external-reproduction checklist, and source audit.'),
        ('provenance/','Unchanged original manuscript and user concept note; historical inputs, not new evidence.')],[2.15,4.6])
    heading('Main reproduction commands')
    for line in ['python -m pip install -r requirements.txt','python -m pytest -q','python scripts/run_experiments.py --output reproduced_results','python scripts/verify_exact.py --results reproduced_results','python scripts/verify_measurements.py --results reproduced_results']:
        p=doc.add_paragraph();r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9)
        p.paragraph_format.space_after=Pt(4)
    para('Use a fresh Python environment, preserve the delivered results, and compare the new output with the archive. The exact tested versions are in requirements-lock.txt and results/environment.json. No GPU, paid data source, cloud credentials, or brokerage connection is required.')
    para('The ZIP manifest detects accidental file changes; it is not a cryptographic signature of scientific authenticity. Authorship, licensing, disclosure, and submission decisions remain with the human researcher.')

    page('6  |  The next gates for a genuinely important paper')
    table(['Gate','Required evidence before advancing'],[
        ('1. External mathematical and novelty review','An independent specialist checks the proof and closest positive-realization/moment literature. Record objections and resolve them explicitly.'),
        ('2. Stronger equal-information comparisons','Implement full-horizon positivity-constrained or moment-feasibility inference; match information, tuning, and compute; retain cases where conventional methods work.'),
        ('3. Feasible measurement protocol','Demonstrate how kernel observations, state resets, units, a, and the learning-rate floor are obtained and audited. Model uncertainty rather than assume it away.'),
        ('4. Financially meaningful mechanism','Add inventory, funding, solvency, liquidity, and strategic behavior. Test whether the hidden-tail mechanism remains measurable and distinct from standard explanations.'),
        ('5. Architecture and error robustness','Vary adaptation families, signed responses, nonlinear impact, model updates, omitted states, and correlated measurement errors.'),
        ('6. Independent external reproduction','A researcher not involved in this build runs the release, reproduces figures, audits assumptions, and reports failures without assistance.'),
        ('7. Empirical and decision value','Use defensible observed or laboratory evidence, compare conventional indicators, and show a decision benefit without offsetting financial harm.')],[2.15,4.6])
    heading('The next single target')
    para('Establish whether a full-horizon, positivity-constrained conventional estimator can resolve the same hidden-feedback cases under exactly the same information. This is the cleanest next comparison because it tests whether the useful contribution is the diagnostic itself, the information requirement, or only the illustrative construction.',True)
    para('The present work is a scientific foothold: a proof, a counterexample, real executed synthetic measurements, and a reproducible package. Paradigm-level influence would require the result to survive those gates and become useful to other researchers. No document length, theorem count, or synthetic accuracy score establishes that outcome by itself.')
    out=ROOT/'docs/Research_Findings_and_Next_Steps.docx'
    doc.save(out);print(out)

if __name__=='__main__':main()
